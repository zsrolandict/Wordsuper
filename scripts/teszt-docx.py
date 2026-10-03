"""
A valódi Wordben végzett kipróbáláshoz: egy szerződéstervezet, amelyben minden új funkció talál valamit.
Futtatás (python-docx 1.2+ kell): python scripts/teszt-docx.py  ->  docs/teszt/WordWriter-teszt.docx
Az ellenőrzőlista: docs/TESZT-WORDBEN.md (minden hibát a szövegével nevez meg).
"""
import copy
import os
from datetime import datetime, timezone

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

OUT = os.path.join(os.path.dirname(__file__), "..", "docs", "teszt", "WordWriter-teszt.docx")

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)

revision_id = [900]


def para(text="", *, bold=False, font=None, size=None, align=None, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        run = p.add_run(text)
        run.bold = bold or None
        if font:
            run.font.name = font
        if size:
            run.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    return p


def run(p, text, **props):
    r = p.add_run(text)
    for key, value in props.items():
        if key == "font":
            r.font.name = value
        elif key == "highlight":
            r.font.highlight_color = value
        elif key == "hidden":
            r.font.hidden = value
        else:
            setattr(r, key, value)
    return r


def tracked(p, text, kind, author="Dr. Ellenfél Ügyvéd"):
    """A pending tracked change from the other side: inserted (w:ins) or deleted (w:del) text"""
    revision_id[0] += 1
    change = OxmlElement("w:ins" if kind == "ins" else "w:del")
    change.set(qn("w:id"), str(revision_id[0]))
    change.set(qn("w:author"), author)
    change.set(qn("w:date"), datetime(2026, 9, 30, 10, 0, tzinfo=timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"))
    r = OxmlElement("w:r")
    t = OxmlElement("w:t" if kind == "ins" else "w:delText")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    change.append(r)
    p._p.append(change)


# ---- Fejléc a meglévő élőfej cseréjének kipróbálásához (Elemek fül: rákérdez, mielőtt felülírja)
doc.sections[0].header.paragraphs[0].text = "TERVEZET – régi élőfej"

# ---- Cím (más betűtípus: a Formázás egységesíti)
para("INGATLAN-ADÁSVÉTELI SZERZŐDÉS", bold=True, font="Arial", size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
para()
p = para(
    "amely létrejött egyrészről az ABC Ingatlanfejlesztő Kft. (székhely: 1051 Budapest, Példa utca 1.; cégjegyzékszám: "
    "01-09-123456; a továbbiakban: „Eladó”), másrészről Kovács János (lakcím: 1111 Budapest, Minta tér 2.; a továbbiakban: "
    "„Vevő”), valamint Nagy Éva (a továbbiakban: „Tulajdonostárs”) között az alulírott helyen és napon, az alábbi feltételekkel:",
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
)
# Két üres sor: a második „fölösleges” (Üres sorok kártya)
para()
para()

# ---- 1. fejezet: álcímsor (félkövér, nem címsorstílus), Calibri
para("1. A SZERZŐDÉS TÁRGYA", bold=True, font="Calibri", size=11)
p = para()
run(p, "1.1 Az Eladó eladja, a Vevő megvásárolja a budapesti 12345/6 hrsz. alatti ingatlant (a továbbiakban: „Ingatlan”).")
doc.add_comment(p.runs, text="Ellenőrizni a hrsz-t a tulajdoni lapon!", author="Dr. Belső Kolléga", initials="BK")
# Hányadok: 1/2 + 1/3 = 5/6, nem 1 – és dupla szóközök
para("1.2  Az Ingatlan  tulajdoni hányadai a szerződés aláírását követően: Vevő 1/2 és Tulajdonostárs 1/3 arányban.")
# Felek elnevezése: a „Vevő” egyes számban definiált, itt többes számban
para("1.3 A Vevők kijelentik, hogy az Ingatlant megtekintették, annak állapotát ismerik.")

# ---- 2. fejezet
para("2. VÉTELÁR", bold=True, font="Calibri", size=11)
p = para()
run(p, "2.1 A vételár ")
run(p, "45.000.000,- Ft, azaz negyvenötmillió forint", highlight=WD_COLOR_INDEX.YELLOW)
run(p, ", amely az Ingatlan teljes ellenértéke.")
# Szám–betű eltérés: 4 500 000 ≠ négymillió-hatszázezer
para("2.2 A foglaló összege 4.500.000 Ft (azaz négymillió-hatszázezer forint), amelyet a Vevő az aláírással egyidejűleg fizet meg.")
p = para()
run(p, "2.3 A vételár fennmaradó részét a ")
run(p, "Vevő", bold=True)  # definiált fogalom félkövéren használva (Definiált fogalmak egységes kiemelése)
run(p, " 2026.10.15-ig - legkésőbb a birtokbaadás napjáig - fizeti meg")
tracked(p, " és járulékait", "ins")
run(p, ".")
p = para()
run(p, "2.4 A késedelmi kamat mértéke a Ptk. szerinti mérték , a fizetési határidő ")
tracked(p, "8", "del")
run(p, "8-15 nap.")
para("2.5 Az Eladó bankszámlaszáma: [●].")

# ---- 4. fejezet: kimaradt a 3. (Számozás-ellenőrzés)
para("4. BIRTOKBAADÁS", bold=True, font="Calibri", size=11)
para("4.1 A birtokbaadás napja: 2026. ___________ hó ___ napja.")
para("4.2 Fontos megjegyezni, hogy a birtokbaadás zökkenőmentes lebonyolítása kulcsfontosságú mindkét fél számára. ✅")
# Markdown-csillagok és egy láthatatlan (nulla szélességű) szóköz
para("4.3 Az **Ingatlan** közüzemi mérőóra-állásait a felek jegyzőkönyvben​ rögzítik.")

# ---- 5. fejezet: valódi Címsor 1 stílus, angolos nagybetűs cím (AI-nyom)
para("5. A Felek Szavatossági Kötelezettségei", style="Heading 1")
para('5.1 Az Eladó szavatol azért, hogy az Ingatlan per-, teher- és igénymentes. Az "Eladó" a Vevővel szemben a jogszabályok szerint felel.')
p = para()
run(
    p,
    "5.2 Tekintettel arra, hogy az Ingatlan a Vevő birtokába kerül, a Vevő viseli a kárveszély átszállásától kezdődően az "
    "Ingatlannal kapcsolatos valamennyi terhet, ideértve különösen a közüzemi díjakat, az adókat és illetékeket, valamint "
    "a társasházi közös költséget, kivéve, ha jelen szerződés másként rendelkezik.",
    bold=True,
)
# Félkövér is, hogy a bekezdés egésze félkövér maradjon (különben a Word vegyesnek látja)
run(p, " BELSŐ MEGJEGYZÉS: az ügyfél max. 40 M Ft-ig menne el.", hidden=True, bold=True)

# ---- Táblázat (a cellák igazítása és térköze)
table = doc.add_table(rows=3, cols=2)
table.style = "Table Grid"
for row, (label, value) in enumerate([("Fél", "Képviselő"), ("Eladó", "Szabó Péter ügyvezető"), ("Vevő", "saját maga")]):
    table.cell(row, 0).text = label
    table.cell(row, 1).text = value

# ---- 6. fejezet
para("6. ZÁRÓ RENDELKEZÉSEK", bold=True, font="Calibri", size=11)
# [Ptk.] nem kitöltetlen hely
para("6.1 A jelen szerződésben nem szabályozott kérdésekben a Polgári Törvénykönyv [Ptk.] rendelkezései az irányadók.")
para(
    "6.2 A felek a szerződést elolvasás és értelmezés után, mint akaratukkal mindenben megegyezőt, jóváhagyólag írják alá.",
    font="Arial",
    size=10.5,
)
para("Kelt: Budapest, 2026. október 3.")
para()
para()
para()
# Aláírási sor: nem kitöltetlen hely
para("______________________")
para("Eladó")

# ---- Melléklet: a számozás itt újraindul, nem hiba
para("1. számú melléklet", bold=True)
para("1. Az Ingatlan alaprajza")
para("2. A mérőórák fényképei")

# ---- Dokumentum-tulajdonságok (Kiküldés előtti ellenőrzés: szerzői adatok)
props = doc.core_properties
props.author = "Teszt Elek (belső)"
props.last_modified_by = "Dr. Belső Kolléga"
props.title = "Tervezet v3"
props.subject = "Kovács-ügy"
props.keywords = "alku, belső"
props.comments = "belső megjegyzés: az ügyfél max. 40 M-ig menne el"
props.category = "Ügyfélanyag"

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print(f"Kész: {os.path.normpath(OUT)}")
